"""Generate the StudyBuddy project charter as a Microsoft Word document (Albanian)."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ORANGE = RGBColor(0xE8, 0x6A, 0x17)
DARK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = "F8F1EA"
HEADER_BG = "1B2A4A"
ACCENT_BG = "E86A17"


def set_run_font(run, name="Calibri", size=11, bold=False, color=DARK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="D6D3D1"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_text(cell, text, *, bold=False, color=DARK, size=10, align="left", fill=None):
    if fill:
        shade_cell(cell, fill)
    set_cell_borders(cell)
    cell.vertical_alignment = 1
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color=NAVY)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "E86A17")
        pBdr.append(bottom)
        pPr.append(pBdr)
    else:
        set_run_font(run, size=13, bold=True, color=ORANGE)
    return p


def add_body(doc, text, *, italic=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=11, color=DARK, italic=italic)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    bullet = p.add_run("•  ")
    set_run_font(bullet, size=11, color=ORANGE, bold=True)
    if bold_prefix:
        b = p.add_run(bold_prefix)
        set_run_font(b, size=11, bold=True, color=NAVY)
        r = p.add_run(text)
        set_run_font(r, size=11, color=DARK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=DARK)
    return p


def add_meta_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, (label, value) in enumerate(rows):
        fill = "F4F6FA" if i % 2 == 0 else "FFFFFF"
        set_cell_text(table.cell(i, 0), label, bold=True, color=NAVY, size=10, fill=fill)
        set_cell_text(table.cell(i, 1), value, color=DARK, size=10, fill=fill)
        table.cell(i, 0).width = Cm(5.4)
        table.cell(i, 1).width = Cm(11.4)
    return table


def add_grid_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        set_cell_text(table.cell(0, j), h, bold=True, color=WHITE, size=10, fill=HEADER_BG)
    for i, row in enumerate(rows):
        fill = ROW_ALT if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row):
            set_cell_text(table.cell(i + 1, j), val, size=10, fill=fill, color=DARK)
    return table


def set_narrow_margins(section):
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)


def set_cover_margins(section):
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)


def set_row_height_cm(row, cm):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(Cm(cm).twips)))
    trHeight.set(qn("w:hRule"), "exact")
    trPr.append(trHeight)


def clear_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_page_start(section, start):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:start"), str(start))


def add_cover_page(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, HEADER_BG)
    clear_cell_borders(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_row_height_cm(table.rows[0], 26.8)

    used_first = False

    def cover_p(text, *, size, bold=False, color=WHITE, space_before=0, space_after=6, italic=False):
        nonlocal used_first
        p = cell.paragraphs[0] if not used_first else cell.add_paragraph()
        used_first = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color, italic=italic)
        return p

    cover_p("STUDYBUDDY", size=14, bold=True, color=ORANGE, space_before=24, space_after=10)
    cover_p("KARTA E PROJEKTIT", size=32, bold=True, color=WHITE, space_before=6, space_after=4)
    cover_p("Project Charter", size=16, italic=True, color=RGBColor(0xCB, 0xD5, 0xE1), space_after=10)
    cover_p("────────────", size=14, color=ORANGE, space_before=4, space_after=16)
    cover_p(
        "Asistent studimi me inteligjencë artificiale",
        size=14,
        color=RGBColor(0xE5, 0xE7, 0xEB),
        space_after=4,
    )
    cover_p("për studentët në Shqipëri", size=14, color=RGBColor(0xE5, 0xE7, 0xEB), space_after=28)
    cover_p("SB-CHARTER-2026-01  ·  Versioni 1.0", size=12, color=RGBColor(0xCB, 0xD5, 0xE1), space_after=4)
    cover_p(date.today().strftime("%d.%m.%Y"), size=12, color=RGBColor(0xCB, 0xD5, 0xE1), space_after=4)
    cover_p("Draft për miratim nga sponsori", size=12, italic=True, color=ORANGE, space_after=36)
    cover_p("studybuddy.al   ·   app.studybuddy.al", size=11, color=RGBColor(0x94, 0xA3, 0xB8), space_after=8)
    cover_p(
        "Konfidencial — për përdorim të brendshëm",
        size=10,
        italic=True,
        color=RGBColor(0x94, 0xA3, 0xB8),
        space_after=12,
    )


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("StudyBuddy  ·  Karta e Projektit  ·  Konfidencial — për përdorim të brendshëm")
    set_run_font(run, size=8, color=MUTED)
    run2 = p.add_run("\t")
    set_run_font(run2, size=8)
    # page numbers
    run3 = p.add_run("Faqja ")
    set_run_font(run3, size=8, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r._r.append(fld1)
    r2 = p.add_run()
    r2._r.append(instr)
    r3 = p.add_run()
    r3._r.append(fld2)
    set_run_font(r, size=8, color=MUTED)
    set_run_font(r2, size=8, color=MUTED)
    set_run_font(r3, size=8, color=MUTED)


def add_header_bar(doc):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, HEADER_BG)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("STUDYBUDDY")
    set_run_font(r, size=11, bold=True, color=ORANGE)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run("Karta e Projektit  ·  Project Charter")
    set_run_font(r2, size=18, bold=True, color=WHITE)
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(10)
    r3 = p3.add_run("Asistent studimi me inteligjencë artificiale për studentët në Shqipëri")
    set_run_font(r3, size=11, color=RGBColor(0xE5, 0xE7, 0xEB))


def build():
    doc = Document()
    cover_section = doc.sections[0]
    set_cover_margins(cover_section)
    cover_section.header.is_linked_to_previous = False
    cover_section.footer.is_linked_to_previous = False
    cover_section.header.paragraphs[0].clear()
    cover_section.footer.paragraphs[0].clear()

    add_cover_page(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_narrow_margins(body_section)
    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False
    body_section.header.paragraphs[0].clear()
    set_page_start(body_section, 1)
    add_footer(body_section)

    add_header_bar(doc)
    doc.add_paragraph()

    add_body(
        doc,
        "Ky dokument autorizon zyrtarisht ekzistencën e projektit StudyBuddy dhe i jep "
        "menaxherit të projektit të drejtën të përdorë burimet e organizatës. Karta "
        "nuk zëvendëson planin e detajuar të projektit; ajo përcakton qëllimin, "
        "kufijtë, suksesin e matshëm dhe qeverisjen në nivel të lartë.",
        italic=True,
        space_after=12,
    )

    add_heading_styled(doc, "1. Të dhënat e dokumentit")
    add_meta_table(
        doc,
        [
            ("Emri i projektit", "StudyBuddy — Asistent studimi me IA"),
            ("Kodi / referenca", "SB-CHARTER-2026-01"),
            ("Versioni", "1.0"),
            ("Data e hartimit", date.today().strftime("%d.%m.%Y")),
            ("Statusi", "Draft për miratim nga sponsori"),
            ("Faqja e marketingut", "https://studybuddy.al"),
            ("Aplikacioni", "https://app.studybuddy.al"),
            ("Tregu fillestar", "Shqipëria (me gjuhë shqipe dhe angleze)"),
            ("Lloji i projektit", "Produkt softuerik SaaS / EdTech"),
        ],
    )

    add_heading_styled(doc, "2. Qëllimi dhe nevoja e biznesit")
    add_body(
        doc,
        "Studentët në Shqipëri dhe në rajon kalojnë orë të tëra duke përpunuar PDF, slajde "
        "dhe shënime, pa mjete të përballueshme që i kthejnë materialet e tyre në përmbledhje, "
        "karta memorie, kuize dhe bisedë të bazuar në dokument. Mjetet ekzistuese globale janë "
        "shpesh të shtrenjta, të palokalizuara ose të përgjithshme — nuk janë ndërtuar rreth "
        "materialit që studenti ka ngarkuar.",
    )
    add_body(
        doc,
        "StudyBuddy i adreson këtë boshllëk: studenti ngarkon dokumentin e vet dhe merr një "
        "komplet studimi të gjeneruar nga IA, plus chat RAG që përgjigjet vetëm nga ai material. "
        "Nëse projekti nuk vazhdon, mbetet një produkt i ndërtuar pa autorizim formal, pa kufij "
        "të qartë të shtrirjes dhe pa kriter suksesi për fazën e tregut.",
    )

    add_heading_styled(doc, "3. Objektivat e matshme (SMART) dhe kriteret e suksesit")
    add_body(doc, "Objektivat e mëposhtme janë të nivelit të kartës; do të detajohen në planin e projektit.")
    add_grid_table(
        doc,
        ["Objektivi", "Treguesi", "Afati i synuar"],
        [
            (
                "Produkt i qëndrueshëm në prodhim",
                "Aplikacioni në app.studybuddy.al dhe faqja studybuddy.al janë të aksesueshme, me autentifikim Clerk dhe rrjedhën Ngarko → Përmbledhje / Karta / Kuiz / Chat",
                "Tashmë (bazë e produktit)",
            ),
            (
                "Përdorues aktivë fillestarë",
                "Të paktën 100 përdorues të regjistruar dhe 50 MAU në tregun shqiptar",
                "6 muaj pas miratimit",
            ),
            (
                "Cilësi e përgjigjeve",
                "Chat-i refuzon pyetjet jashtë temës dhe përgjigjet në gjuhën e zgjedhur (SQ/EN) mbi materialin e ngarkuar",
                "Të vazhdueshme",
            ),
            (
                "Qëndrueshmëri kostoje",
                "Kosto infrastrukture ≤ ~3 USD/muaj në shkallë të vogël (10 përdorues) dhe e dokumentuar për shkallë mesatare/të madhe",
                "Të vazhdueshme",
            ),
            (
                "Gati për të ardhura",
                "Modeli i pagesës (p.sh. Stripe / Lemon Squeezy) i specifikuar dhe i vendosur për lansim; pa të ardhura të rreme në komunikim publik",
                "12 muaj pas miratimit",
            ),
        ],
    )

    add_heading_styled(doc, "4. Kërkesat e nivelit të lartë")
    add_bullet(doc, "Ngarkim i materialeve të studimit (PDF, DOCX, PPTX, TXT, imazhe/HEIC) me kufij mujorë.", bold_prefix="Ngarkim: ")
    add_bullet(doc, "Përmbledhje, karta memorie dhe kuize të gjeneruara nga IA, me mundësi fokusi tematik.", bold_prefix="Gjenerim: ")
    add_bullet(doc, "Chat me RAG (Pinecone + OpenAI) i lidhur me dokumentin ose grupin e dokumenteve.", bold_prefix="Chat: ")
    add_bullet(doc, "Shënime të integruara (redaktues, dosje) për të ruajtur mësimin pranë materialit.", bold_prefix="Shënime: ")
    add_bullet(doc, "Hyrje me Clerk; ndërfaqe dygjuhëshe (shqip / anglisht).", bold_prefix="Akses: ")
    add_bullet(doc, "Paneli i superpërdoruesit për kufij ngarkimi/chat dhe menaxhim përdoruesish.", bold_prefix="Administrim: ")
    add_bullet(doc, "Prodhimi i dyfishtë: marketing (studybuddy.al) dhe aplikacion (app.studybuddy.al).", bold_prefix="Platformë: ")

    add_heading_styled(doc, "5. Shtrirja: çfarë hyn dhe çfarë nuk hyn")
    add_heading_styled(doc, "5.1 Në shtrirje", level=2)
    for item in [
        "Aplikacioni web i studentit dhe paneli administrativ.",
        "Faqja e marketingut dhe identiteti i domenit .al.",
        "Procesimi i dokumenteve, embeddings, RAG dhe kufijtë e përdorimit.",
        "Lokalizimi SQ/EN i ndërfaqes.",
        "Analiza e kostove të infrastrukturës dhe përgatitja e modelit të të ardhurave.",
        "Përmirësime të produktit ekzistues (qëndrueshmëri, kufij, përvoja e përdoruesit, fshirje llogarie nga admin).",
    ]:
        add_bullet(doc, item)

    add_heading_styled(doc, "5.2 Jashtë shtrirjes (në këtë fazë)", level=2)
    for item in [
        "Aplikacion native iOS / Android (përveçse si vendim i mëvonshëm i sponsorit).",
        "LMS i plotë (notat e shkollës, orari i klasës, detyrat e mësuesit).",
        "Integrime me Google Classroom, Microsoft Teams ose Moodle.",
        "Model i vetë-trajnuar LLM; përdoret OpenAI si furnizues.",
        "Zgjerim gjeografik përtej Shqipërisë si treg primar (mund të rihapet).",
        "Mbështetje e dedikuar 24/7 dhe SaaS i veçantë për email marketing.",
    ]:
        add_bullet(doc, item)

    add_heading_styled(doc, "6. Produkte kryesore (deliverables)")
    add_grid_table(
        doc,
        ["Produkti", "Përshkrimi", "Pranimi"],
        [
            ("Aplikacioni StudyBuddy", "Web app i gatshëm për studentë në app.studybuddy.al", "Sponsori konfirmon rrjedhën kryesore të përdorimit"),
            ("Faqja e marketingut", "studybuddy.al me CTA drejt regjistrimit", "Domeni dhe CTA funksionojnë"),
            ("Karta e projektit", "Ky dokument, i nënshkruar", "Nënshkrimi i sponsorit"),
            ("Dokumentimi i kostove", "Analiza e infrastrukturës (docs/COST_ANALYSIS.md)", "Rishikim i sponsorit"),
            ("Plan pagesash", "Specifikim i niveleve dhe integrimit të faturimit", "Vendim i sponsorit për lansim"),
        ],
    )

    add_heading_styled(doc, "7. Afati i përmbledhur (milestones)")
    add_body(
        doc,
        "Datat e mëposhtme janë objektiva të nivelit të lartë, jo zotime kontraktuale. "
        "Plani i detajuar do t’i rishikojë pas miratimit të kartës.",
    )
    add_grid_table(
        doc,
        ["Faza", "Përshkrimi", "Objektivi"],
        [
            ("F0 — Baza", "Produkt live: ngarkim, gjenerim, chat, shënime, admin, i18n", "E përfunduar / në prodhim"),
            ("F1 — Stabilizim", "Qëndrueshmëri (kufij chat/ngarkim, fshirje përdoruesi, gabime 500), matje përdoruesish", "0–3 muaj"),
            ("F2 — Tregu", "Përvetësim në Shqipëri, feedback studentësh, përmirësim i onboarding-ut", "3–6 muaj"),
            ("F3 — Të ardhura", "Lansim i modeleve të pagesës pa pretendime të rreme të të ardhurave", "6–12 muaj"),
            ("Mbyllja e fazës", "Rishikim i kartës: vazhdim, rishtrirje ose mbyllje", "12 muaj"),
        ],
    )

    add_heading_styled(doc, "8. Buxheti dhe burimet e para-aprovuara")
    add_body(
        doc,
        "Buxheti i infrastrukturës është i ulët në fazën aktuale. Sipas analizës së kostove "
        "(qershor 2026, USD): rreth 2–3 USD/muaj me ~10 përdorues (kryesisht plane falas + domen); "
        "rreth 48–98 USD/muaj në skenarin ~1.000 të regjistruar / ~400 MAU, kryesisht nga OpenAI "
        "dhe mundësisht Supabase Pro / Vercel Pro. Pagat, kontraktorët dhe mjetet e mbështetjes "
        "së klientit nuk janë pjesë e këtij zarfi infrastrukture.",
    )
    add_bullet(doc, "Hosting: Vercel (marketing + app).", bold_prefix="Stack: ")
    add_bullet(doc, "Autentifikim: Clerk.", bold_prefix="Stack: ")
    add_bullet(doc, "Bazë të dhënash dhe skedarë: Supabase.", bold_prefix="Stack: ")
    add_bullet(doc, "Vektorë RAG: Pinecone; LLM/embeddings: OpenAI (gpt-4o-mini, text-embedding-3-small).", bold_prefix="Stack: ")
    add_body(
        doc,
        "Toleranca: rritje e kostos mujore të infrastrukturës mbi 100 USD kërkon rishikim të sponsorit. "
        "Çdo angazhim pagese ndaj palëve të treta mbi këtë prag kërkon miratim paraprak.",
    )

    add_heading_styled(doc, "9. Palët e interesuara dhe ekipi")
    add_grid_table(
        doc,
        ["Roli", "Emri / njësia", "Përgjegjësia"],
        [
            ("Sponsor i projektit", "[Emri i sponsorit]", "Autorizon kartën, buxhetin dhe ndryshimet e shtrirjes"),
            ("Menaxher i projektit / themelues", "[Emri]", "Drejton ekzekutimin dhe përdor burimet e autorizuara"),
            ("Zhvillim produkti", "Ekipi StudyBuddy", "Ndërton, mban dhe vendos aplikacionin"),
            ("Klient / përdorues", "Studentë në Shqipëri", "Përdorin produktin; japin feedback"),
            ("Furnizues kyç", "OpenAI, Supabase, Clerk, Pinecone, Vercel", "Shërbime të palëve të treta; rrezik varësie"),
        ],
    )
    add_body(
        doc,
        "Plotësoni emrat në kllapa para nënshkrimit. Menaxheri i projektit ka autoritetin të "
        "vendosë për detaje zbatimi brenda shtrirjes dhe buxhetit; ndryshimet e shtrirjes, "
        "tregut primar ose modelit të të ardhurave i kthehen sponsorit.",
        italic=True,
    )

    add_heading_styled(doc, "10. Supozime dhe kufizime")
    add_heading_styled(doc, "10.1 Supozime", level=2)
    for item in [
        "Studentët kanë akses interneti dhe pajisje web; produkti është web-first.",
        "OpenAI, Supabase, Clerk dhe Pinecone mbeten të përdorshëm me çmime të qëndrueshme.",
        "Domenet studybuddy.al dhe app.studybuddy.al mbeten nën kontrollin e projektit.",
        "Tregu fillestar është Shqipëria; kërkesa ekziston për mjete studimi me IA në shqip.",
        "Faturimi nuk është lansuar ende; komunikimi publik nuk pretendon të ardhura aktuale.",
    ]:
        add_bullet(doc, item)

    add_heading_styled(doc, "10.2 Kufizime", level=2)
    for item in [
        "Buxhet i vogël infrastrukture; pa ekip të madh inxhinierik të palës së tretë.",
        "Varësi nga API të palëve të treta (IA, auth, hosting, vektorë).",
        "Kufij mujorë ngarkimi dhe chat-i për të kontrolluar koston dhe abuzimin.",
        "Nuk ka LMS të plotë as aplikacion native në këtë fazë.",
        "Pajtueshmëria me mbrojtjen e të dhënave (përfshirë fshirjen e llogarisë) duhet respektuar.",
    ]:
        add_bullet(doc, item)

    add_heading_styled(doc, "11. Rreziqet e nivelit të lartë")
    add_grid_table(
        doc,
        ["Rreziku", "Ndikimi", "Përgjigja fillestare"],
        [
            ("Rritje e çmimit ose kufizim i OpenAI", "Kosto e papritur ose ndërprerje e gjenerimit/chat", "Kufij përdorimi; monitorim kostoje; plan B i modelit"),
            ("Gabime 500 / dështim i kufijve të chat-it", "Humbje besimi i përdoruesit", "Stabilizim i RPC/migrimeve; testim para lansimit të faturimit"),
            ("Përdorim jashtë temës ose abuzim i IA", "Kosto dhe cilësi e dobët", "Filtrim relevancë; kufij mujorë"),
            ("Vonim i modelit të pagesës", "Mungesë qëndrueshmërie financiare", "Mbajtje kostoje të ulët deri në lansim"),
            ("Varësi nga një person kyç", "Ndalim i zhvillimit", "Dokumentim, depo Git, kartë e nënshkruar"),
            ("Çështje private / fshirje llogarie", "Rrezik ligjor dhe reputacioni", "Rrjedha e fshirjes nga llogaria dhe nga admin"),
        ],
    )

    add_heading_styled(doc, "12. Kërkesat e miratimit dhe kriteret e daljes")
    add_heading_styled(doc, "12.1 Miratimi", level=2)
    add_body(
        doc,
        "Karta hyn në fuqi me nënshkrimin e sponsorit. Ndryshimet materiale (shtrirje, treg, "
        "buxhet mbi tolerancë, vendim për faturim) kërkojnë rishikim të kartës. Produktet "
        "pranohen sipas kolonës «Pranimi» në seksionin 6.",
    )
    add_heading_styled(doc, "12.2 Dalja / mbyllja", level=2)
    add_bullet(doc, "Objektivat F0–F3 janë përmbushur ose janë rinegociuar me shkrim.", bold_prefix="Sukses: ")
    add_bullet(
        doc,
        "Sponsori vendos ndërprerje nëse kostoja tejkalon tolerancën pa plan, nëse rreziku i palëve të treta bëhet i papranueshëm, ose nëse tregu nuk vërteton kërkesën brenda 12 muajve.",
        bold_prefix="Ndërprerje e parakohshme: ",
    )

    add_heading_styled(doc, "13. Autorizimi")
    add_body(
        doc,
        "Duke nënshkruar, sponsori autorizon ekzistencën e projektit StudyBuddy dhe i jep "
        "menaxherit të projektit autoritetin të zbatojë këtë kartë, të përdorë burimet e "
        "përshkruara dhe të raportojë te sponsori për përparimin, rreziqet dhe kërkesat e ndryshimit.",
    )

    sign = doc.add_table(rows=4, cols=3)
    headers = ["Roli", "Emri / firma", "Data"]
    values = [
        ["Sponsor i projektit", "________________________", "____ / ____ / ________"],
        ["Menaxher i projektit", "________________________", "____ / ____ / ________"],
        ["Palë tjetër (opsionale)", "________________________", "____ / ____ / ________"],
    ]
    for j, h in enumerate(headers):
        set_cell_text(sign.cell(0, j), h, bold=True, color=WHITE, size=10, fill=ACCENT_BG)
    for i, row in enumerate(values):
        fill = "FFFFFF" if i % 2 == 0 else "F8F1EA"
        for j, val in enumerate(row):
            set_cell_text(sign.cell(i + 1, j), val, size=10, fill=fill)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(14)
    r = note.add_run(
        "Referenca: PMBOK / PMI — karta e projektit autorizon punën dhe palën që e drejton atë. "
        "Ky dokument është hartuar në shqip për StudyBuddy (studybuddy.al / app.studybuddy.al)."
    )
    set_run_font(r, size=9, italic=True, color=MUTED)

    out = Path(__file__).resolve().parent / "Karta_e_Projektit_StudyBuddy.docx"
    doc.save(out)
    print(str(out))


if __name__ == "__main__":
    build()
